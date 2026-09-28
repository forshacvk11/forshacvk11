"""
Tender Scout — a LangGraph workflow that finds, screens and drafts EOIs for
consultancy tenders that match your practice.

    fetch ──► dedupe ──► screen ──►(any strong matches?)──► draft_eoi ──► update_tracker ──► report
                                          └──────────── no ──────────────────────────────────┘

Run once:      python tender_agent.py
Test offline:  python tender_agent.py --demo      (no internet or API key needed)
Schedule it daily with Windows Task Scheduler or cron (see README.md).
"""
import os, sys, json, re, sqlite3, hashlib, datetime as dt, argparse
from pathlib import Path
from typing import TypedDict

from langgraph.graph import StateGraph, START, END

HERE = Path(__file__).parent
CONFIG = json.loads((HERE / "config.json").read_text(encoding="utf-8"))
PROFILE = (HERE / "firm_profile.md").read_text(encoding="utf-8")
OUT = HERE / CONFIG.get("output_folder", "output"); OUT.mkdir(exist_ok=True)
DB = HERE / "seen_tenders.db"
DEMO = False


# ------------------------------------------------------------------ STATE
class State(TypedDict, total=False):
    raw: list          # everything fetched
    new: list          # not seen before
    screened: list     # with score + reasons
    strong: list       # score >= threshold
    drafts: list       # file paths of EOI drafts
    report: str


# ------------------------------------------------------------------ LLM
def llm(model_key: str):
    from langchain_anthropic import ChatAnthropic
    return ChatAnthropic(model=CONFIG["models"][model_key], max_tokens=4000, temperature=0)


def ask_json(model_key, system, user):
    msg = llm(model_key).invoke([("system", system), ("user", user)])
    text = msg.content if isinstance(msg.content, str) else "".join(
        b.get("text", "") for b in msg.content if isinstance(b, dict))
    text = re.sub(r"^```(json)?|```$", "", text.strip(), flags=re.M).strip()
    return json.loads(text)


# ------------------------------------------------------------------ NODE 1: FETCH
def fetch(state: State) -> State:
    if DEMO:
        return {"raw": DEMO_TENDERS}
    import requests
    from bs4 import BeautifulSoup
    items, headers = [], {"User-Agent": "Mozilla/5.0 (TenderScout)"}

    # A) World Bank procurement notices for Nigeria (public API)
    try:
        r = requests.get("https://search.worldbank.org/api/v2/procnotices",
                         params={"format": "json", "countrycode_exact": "NG", "rows": 60,
                                 "srt": "noticedate", "order": "desc"}, headers=headers, timeout=30)
        notices = r.json().get("procnotices", {})
        notices = notices.values() if isinstance(notices, dict) else notices
        for n in notices:
            title = n.get("bid_description") or n.get("project_name") or ""
            items.append({
                "source": "World Bank",
                "title": title.strip(),
                "client": n.get("contact_organization", "") or n.get("borrower", ""),
                "type": n.get("notice_type", ""),
                "method": n.get("procurement_method_name", ""),
                "deadline": (n.get("submission_deadline_date") or n.get("submission_date") or "")[:10],
                "url": f"https://projects.worldbank.org/en/projects-operations/procurement-detail/{n.get('id','')}",
                "text": (n.get("notice_text") or title)[:3000],
            })
    except Exception as e:
        print("World Bank fetch failed:", e)

    # B) Tender websites from config.json: collect links that look like consultancy calls
    pattern = re.compile(CONFIG["link_keywords"], re.I)
    for site in CONFIG["tender_pages"]:
        try:
            html = requests.get(site, headers=headers, timeout=30).text
            soup = BeautifulSoup(html, "html.parser")
            for a in soup.find_all("a", href=True):
                t = " ".join(a.get_text(" ").split())
                if len(t) > 25 and pattern.search(t):
                    items.append({"source": site, "title": t, "client": "", "type": "", "method": "",
                                  "deadline": "", "url": requests.compat.urljoin(site, a["href"]), "text": t})
        except Exception as e:
            print("Fetch failed:", site, e)
    print(f"fetched {len(items)} items")
    return {"raw": items}


# ------------------------------------------------------------------ NODE 2: DEDUPE
def dedupe(state: State) -> State:
    con = sqlite3.connect(DB)
    con.execute("CREATE TABLE IF NOT EXISTS seen(id TEXT PRIMARY KEY, title TEXT, first_seen TEXT)")
    new = []
    for it in state["raw"]:
        key = hashlib.sha1((it["title"].lower() + it["url"]).encode()).hexdigest()
        if not con.execute("SELECT 1 FROM seen WHERE id=?", (key,)).fetchone():
            it["id"] = key; new.append(it)
            if not DEMO:
                con.execute("INSERT INTO seen VALUES(?,?,?)", (key, it["title"][:300], dt.date.today().isoformat()))
    con.commit(); con.close()
    print(f"{len(new)} new")
    return {"new": new}


# ------------------------------------------------------------------ NODE 3: SCREEN
SCREEN_SYSTEM = """You screen procurement notices for a small Nigerian civil engineering consultancy.
Score each notice 0-10 for fit with the firm profile below. High scores need ALL of:
consultancy services (not works/goods), a match with the firm's specialisms, and a location
the firm can serve (or any location if it could join as a sub-consultant/key expert).
Return ONLY a JSON list: [{"i": index, "score": int, "role": "Lead|JV member|Sub-consultant|Key expert|Skip",
"why": "one sentence", "sector": "short tag"}]

FIRM PROFILE:
""" + PROFILE


def screen(state: State) -> State:
    items = state["new"]
    if not items:
        return {"screened": [], "strong": []}
    scored = []
    for start in range(0, len(items), 15):          # batches keep prompts small and cheap
        batch = items[start:start + 15]
        listing = "\n".join(f"[{i}] {t['title']} | client: {t['client']} | type: {t['type']} | "
                            f"deadline: {t['deadline']} | {t['text'][:500]}" for i, t in enumerate(batch))
        res = DEMO_SCORES[start:start + 15] if DEMO else ask_json("screen", SCREEN_SYSTEM, listing)
        for r in res:
            t = dict(batch[r["i"]]); t.update(score=r["score"], role=r["role"], why=r["why"], sector=r.get("sector", ""))
            scored.append(t)
    thr = CONFIG["draft_threshold"]
    strong = [t for t in scored if t["score"] >= thr]
    print(f"{len(strong)} strong matches (score >= {thr})")
    return {"screened": scored, "strong": strong}


def route(state: State) -> str:
    return "draft_eoi" if state.get("strong") else "report"


# ------------------------------------------------------------------ NODE 4: DRAFT EOI
DRAFT_SYSTEM = """You write Expressions of Interest for World Bank / AfDB / Nigerian federal consultancy calls.
Use ONLY facts from the firm profile; where a fact is missing write a [PLACEHOLDER] in square brackets.
Never invent past projects, registration numbers or staff. Structure:
1. Cover letter  2. Firm profile & core business  3. Relevant experience (table placeholders if needed)
4. Technical & managerial capability incl. key experts  5. Proposed approach for THIS assignment (specific)
6. Association/JV statement if the role is not Lead  7. Attachments checklist.
Write in clear professional British English. Output Markdown only.

FIRM PROFILE:
""" + PROFILE


def draft_eoi(state: State) -> State:
    paths = []
    for t in state["strong"]:
        user = (f"Assignment: {t['title']}\nClient: {t['client']}\nMethod: {t['method']}\n"
                f"Deadline: {t['deadline']}\nOur likely role: {t['role']}\nLink: {t['url']}\n\nNotice text:\n{t['text']}")
        if DEMO:
            md = f"# EOI — {t['title']}\n\n(demo draft)\n"
        else:
            m = llm("draft").invoke([("system", DRAFT_SYSTEM), ("user", user)])
            md = m.content if isinstance(m.content, str) else "".join(b.get("text", "") for b in m.content if isinstance(b, dict))
        slug = re.sub(r"[^A-Za-z0-9]+", "_", t["title"])[:60].strip("_")
        p = OUT / f"EOI_{dt.date.today():%Y%m%d}_{slug}.md"
        p.write_text(md, encoding="utf-8"); paths.append(str(p))
    return {"drafts": paths}


# ------------------------------------------------------------------ NODE 5: UPDATE TRACKER
def update_tracker(state: State) -> State:
    path = HERE / CONFIG["tracker_path"]
    if not path.exists():
        print("Tracker not found, skipping:", path); return {}
    from openpyxl import load_workbook
    wb = load_workbook(path); ws = wb["Pipeline"]
    row = next((r for r in range(5, 45) if not ws.cell(r, 2).value), None)
    for t in state["strong"]:
        if row is None:
            print("Tracker full — add rows to the Pipeline sheet"); break
        vals = {2: t["title"][:250], 3: t["client"], 4: t["source"], 6: "EOI", 7: t["method"][:20],
                8: t["role"] if t["role"] != "Skip" else None, 14: "Reviewing",
                15: round(t["score"] / 20, 2), 17: "Review AI draft EOI: " + t["why"], 18: t["url"]}
        for c, v in vals.items():
            ws.cell(row, c).value = v
        try:
            ws.cell(row, 11).value = dt.date.fromisoformat(t["deadline"][:10])
        except Exception:
            pass
        row = next((r for r in range(row + 1, 45) if not ws.cell(r, 2).value), None)
    wb.save(path)
    return {}


# ------------------------------------------------------------------ NODE 6: REPORT
def report(state: State) -> State:
    s = sorted(state.get("screened", []), key=lambda t: -t["score"])
    lines = [f"# Tender Scout — {dt.date.today():%d %b %Y}",
             f"New notices: {len(state.get('new', []))}  |  Strong matches: {len(state.get('strong', []))}", ""]
    for t in s:
        if t["score"] < CONFIG["report_min_score"]:
            continue
        lines += [f"## [{t['score']}/10] {t['title']}", f"- Role: {t['role']}  |  Deadline: {t['deadline'] or 'see notice'}",
                  f"- Why: {t['why']}", f"- Link: {t['url']}", ""]
    if state.get("drafts"):
        lines += ["## EOI drafts ready for your review"] + [f"- {Path(p).name}" for p in state["drafts"]]
    text = "\n".join(lines)
    (OUT / f"digest_{dt.date.today():%Y%m%d}.md").write_text(text, encoding="utf-8")
    send_email(text)
    print(text)
    return {"report": text}


def send_email(text):
    e = CONFIG.get("email", {})
    if DEMO or not e.get("enabled"):
        return
    import smtplib
    from email.message import EmailMessage
    msg = EmailMessage(); msg["Subject"] = f"Tender Scout — {dt.date.today():%d %b}"
    msg["From"], msg["To"] = e["from"], e["to"]; msg.set_content(text)
    with smtplib.SMTP_SSL(e["smtp_host"], e.get("smtp_port", 465)) as s:
        s.login(e["from"], os.environ["TENDER_SCOUT_EMAIL_PASSWORD"]); s.send_message(msg)


# ------------------------------------------------------------------ GRAPH
def build_graph():
    g = StateGraph(State)
    for name, fn in [("fetch", fetch), ("dedupe", dedupe), ("screen", screen),
                     ("draft_eoi", draft_eoi), ("update_tracker", update_tracker), ("report", report)]:
        g.add_node(name, fn)
    g.add_edge(START, "fetch"); g.add_edge("fetch", "dedupe"); g.add_edge("dedupe", "screen")
    g.add_conditional_edges("screen", route, {"draft_eoi": "draft_eoi", "report": "report"})
    g.add_edge("draft_eoi", "update_tracker"); g.add_edge("update_tracker", "report"); g.add_edge("report", END)
    return g.compile()


# ------------------------------------------------------------------ DEMO DATA
DEMO_TENDERS = [
    {"source": "demo", "title": "SPIN: Consultancy for construction supervision of Adani irrigation scheme, Enugu State",
     "client": "FMWRS", "type": "REOI", "method": "QCBS", "deadline": "2026-10-20", "url": "https://example.org/1", "text": "Supervision of design & build irrigation works."},
    {"source": "demo", "title": "Supply of office furniture to State Ministry of Health",
     "client": "MoH", "type": "ITB", "method": "NCB", "deadline": "2026-10-10", "url": "https://example.org/2", "text": "Goods."},
    {"source": "demo", "title": "Design of urban storm drainage master plan, Abuja satellite towns",
     "client": "FCDA", "type": "REOI", "method": "CQS", "deadline": "2026-10-25", "url": "https://example.org/3", "text": "Hydrology and drainage design."},
]
DEMO_SCORES = [
    {"i": 0, "score": 8, "role": "Key expert", "why": "Irrigation supervision in the South-East matches hydraulics + CM.", "sector": "irrigation"},
    {"i": 1, "score": 0, "role": "Skip", "why": "Goods supply, not consultancy.", "sector": "goods"},
    {"i": 2, "score": 9, "role": "Lead", "why": "CQS drainage design in FCT is winnable alone.", "sector": "drainage"},
]

if __name__ == "__main__":
    ap = argparse.ArgumentParser(); ap.add_argument("--demo", action="store_true")
    DEMO = ap.parse_args().demo
    if not DEMO and not os.environ.get("ANTHROPIC_API_KEY"):
        sys.exit("Set ANTHROPIC_API_KEY first (see README.md), or run with --demo")
    build_graph().invoke({})